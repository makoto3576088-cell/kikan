#!/usr/bin/env python3
"""出席管理シート（Google Sheets）の「コマ」タブから、植田のコマだけを
kikan の events 表に同期するための SQL を生成する。

使い方:
  python3 tools/sync_sheet.py <入力> [--today YYYY-MM-DD] [--out out.sql]

<入力> は次のどちらか:
  - Google Drive コネクタ download_file_content (exportMimeType=xlsx) の結果を
    保存した JSON ファイル（{content: <base64>, ...}）
  - .xlsx ファイルそのもの

出力は SQL 1 文。Supabase の execute_sql でそのまま実行する。
実行結果は inserted / updated / deleted の件数が 1 行返る。
"""
import argparse
import base64
import collections
import datetime as dt
import io
import json
import re
import sys

import openpyxl

USER_ID = "23d69b8c-8467-455d-99c9-c27bf0f619fa"
TEACHER_NAME = "植田"          # 講師タブの氏名。該当する講師IDをすべて拾う
SKIP_STATES = {"取消", "休講"}  # 予定に出さない状態
WINDOW_BACK, WINDOW_FWD = 7, 60  # 同期する範囲（今日の前後、日数）
WEEK = "月火水木金土日"


def load_workbook(path):
    if path.lower().endswith(".xlsx"):
        return openpyxl.load_workbook(path, data_only=True)
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    raw = base64.b64decode(d["content"])
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=True)


def rows_after_header(ws, first_col_header):
    """ヘッダー行（A列が first_col_header）の次からの行を返す。"""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and r[0] == first_col_header:
            return rows[i + 1:]
    return []


def sid(v):
    if v is None:
        return ""
    s = str(v)
    return s[:-2] if s.endswith(".0") else s


def parse_band(s):
    """'9:00〜9:40' -> ('09:00', '09:40')"""
    m = re.match(r"\s*(\d{1,2}):(\d{2})\s*[〜~～\-–]\s*(\d{1,2}):(\d{2})", str(s or ""))
    if not m:
        return None
    return (f"{int(m[1]):02d}:{m[2]}", f"{int(m[3]):02d}:{m[4]}")


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def campus_of(v):
    s = str(v or "")
    if s.startswith("オーチャード"):
        return "オーチャード校"
    if s.startswith("クレメン"):
        return "クレメンティ校"
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--today", default=None, help="YYYY-MM-DD（省略時はシンガポール時間の今日）")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    if a.today:
        today = dt.date.fromisoformat(a.today)
    else:
        today = (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=8)).date()
    d_from, d_to = today - dt.timedelta(days=WINDOW_BACK), today + dt.timedelta(days=WINDOW_FWD)

    wb = load_workbook(a.input)

    # --- 講師: 氏名が TEACHER_NAME の講師ID
    teacher_ids = {r[0] for r in rows_after_header(wb["講師"], "講師ID") if r[0] and r[1] == TEACHER_NAME}
    if not teacher_ids:
        sys.exit(f"講師タブに {TEACHER_NAME} が見つかりません")

    # --- 生徒: ID -> (氏名, 分)
    students = {}
    for r in rows_after_header(wb["生徒"], "生徒ID"):
        if r[0]:
            students[sid(r[0])] = (r[2] or r[1] or "", r[5])

    # --- 設定: (校舎, 講) -> {'前': (s,e), '後': (s,e)}
    bands = {}
    for r in wb["設定"].iter_rows(values_only=True):
        cells = list(r)
        for i in range(len(cells) - 3):
            if cells[i] in ("オーチャード校", "クレメンティ校") and isinstance(cells[i + 1], str) and cells[i + 1].endswith(("講", "他")):
                b = {}
                if parse_band(cells[i + 2]):
                    b["前"] = parse_band(cells[i + 2])
                if parse_band(cells[i + 3]):
                    b["後"] = parse_band(cells[i + 3])
                if b:
                    bands[(cells[i], cells[i + 1])] = b
    if not bands:
        sys.exit("設定タブの時間帯が読めません")

    # --- 枠: 曜日チェック用（xlsx 書き出しで日付が 1 日ずれることがあるので補正量を推定）
    waku = {}
    for r in rows_after_header(wb["枠"], "枠ID"):
        if r[0]:
            waku[(sid(r[1]), r[6], r[7], r[9])] = r[5]

    koma = [r for r in rows_after_header(wb["コマ"], "コマID") if r[0] and r[1]]
    votes = collections.Counter()
    for r in koma:
        if r[9] != "通常":
            continue
        k = (sid(r[7]), r[3], r[4], r[6])
        if k in waku and waku[k] in WEEK:
            for shift in (-1, 0, 1):
                d = r[1].date() + dt.timedelta(days=shift)
                if WEEK[d.weekday()] == waku[k]:
                    votes[shift] += 1
    shift = votes.most_common(1)[0][0] if votes else 0
    if votes and votes.most_common(1)[0][1] < 0.8 * sum(votes.values()):
        sys.exit(f"日付のずれを特定できません: {dict(votes)}")

    # --- 植田のコマを集める（1:2 は同じ時間枠にまとめる）
    groups = collections.defaultdict(list)  # (date, 校舎, 講, 前後) -> rows
    for r in koma:
        if r[6] not in teacher_ids or r[10] in SKIP_STATES:
            continue
        d = r[1].date() + dt.timedelta(days=shift)
        if not (d_from <= d <= d_to):
            continue
        groups[(d, campus_of(r[2]), r[3], r[4])].append(r)

    events = []
    for (d, campus, ko, half), rs in sorted(groups.items()):
        band = bands.get((campus, ko)) or {}
        if half not in band:
            print(f"時間帯なし: {d} {campus} {ko} {half}", file=sys.stderr)
            continue
        start, end = band[half]
        names, notes, is_absent = [], [], []
        for r in sorted(rs, key=lambda r: (r[5] or 0)):
            s_id = sid(r[7])
            name, mins = students.get(s_id, (None, None))
            if not name:
                name = r[12] or r[8] or (r[9] if r[9] in ("体験", "追加", "面談") else "") or "（名前なし）"
            lab = []
            if r[9] and r[9] != "通常":
                lab.append(r[9])
            if r[10] == "欠席":
                lab.append("欠席")
            names.append(name + (f"（{'・'.join(lab)}）" if lab else ""))
            is_absent.append(r[10] == "欠席")
            if mins == 90 and half == "前" and "後" in band:
                end = band["後"][1]
            if r[16]:
                notes.append(str(r[16]))
            if r[13]:
                notes.append(str(r[13]))
        mins_set = {students.get(sid(r[7]), (None, None))[1] for r in rs} - {None}
        memo = f"{ko} {half}" + (" 1:2" if len(rs) > 1 else "")
        if len(mins_set) == 1:
            memo += f" · {int(mins_set.pop())}分"
        if notes:
            memo += " · " + " / ".join(dict.fromkeys(notes))
        refs = "+".join(sorted(r[0] for r in rs))
        events.append({
            "ext_ref": "sheet:" + refs,
            "date": d.isoformat(),
            "start": start,
            "end": end,
            "title": "・".join(names),
            "location": campus,
            "memo": memo,
        })

    events.sort(key=lambda e: (e["date"], e["start"], e["title"]))

    # --- SQL
    uid = q(USER_ID)
    if events:
        vals = ",\n".join(
            f"({q(e['ext_ref'])},{q(e['date'])},{q(e['start'])},{q(e['end'])},{q(e['title'])},{q(e['location'])},{q(e['memo'])})"
            for e in events
        )
        sql = f"""with src(ext_ref, date, start_time, end_time, title, location, memo) as (values
{vals}
), up as (
  insert into public.events (user_id, title, scope, date, all_day, start_time, end_time, location, memo, repeat, repeat_days, ext_ref)
  select {uid}, title, 'work', date::date, false, start_time::time, end_time::time, location, memo, 'none', '{{}}', ext_ref from src
  on conflict (user_id, ext_ref) where ext_ref is not null do update set
    title = excluded.title, date = excluded.date, start_time = excluded.start_time, end_time = excluded.end_time,
    location = excluded.location, memo = excluded.memo, updated_at = now()
  returning (xmax = 0) as inserted
), del as (
  delete from public.events
  where user_id = {uid} and ext_ref like 'sheet:%' and date between {q(d_from)} and {q(d_to)}
    and ext_ref not in (select ext_ref from src)
  returning id
)
select (select count(*) from up where inserted) as inserted,
       (select count(*) from up where not inserted) as updated,
       (select count(*) from del) as deleted;"""
    else:
        sql = f"""with del as (
  delete from public.events
  where user_id = {uid} and ext_ref like 'sheet:%' and date between {q(d_from)} and {q(d_to)}
  returning id
) select 0 as inserted, 0 as updated, (select count(*) from del) as deleted;"""

    summary = f"-- {len(events)} 件 / {d_from}〜{d_to} / 日付補正 {shift:+d} 日 / 講師ID {sorted(teacher_ids)}\n"
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(sql)
        print(summary.strip())
        for e in events:
            print(f"{e['date']} {e['start']}-{e['end']} {e['title']} @{e['location']}  [{e['memo']}]")
    else:
        print(summary + sql)


if __name__ == "__main__":
    main()
